"""Locators refer to immutable canonical Unicode code points, never UTF-16."""

import hashlib
import zipfile

import pytest


def source(tmp_path, text, kind="txt"):
    from app.rag.contracts import SourceInput

    path = tmp_path / ("source." + kind)
    path.write_bytes(text.encode("utf-8"))
    return SourceInput(
        owner_id=7,
        namespace="production",
        doc_id="d1",
        document_version_id="v1",
        file_path=str(path),
        file_type=kind,
        title="资料",
    )


def test_text_emoji_offsets_and_paragraphs_are_codepoints(tmp_path):
    from app.rag.ingestion import load_canonical_document

    doc = load_canonical_document(
        source(tmp_path, "光合作用需要光。\r\n\r\n😀第二段。")
    )
    assert doc.text == "光合作用需要光。\n\n😀第二段。"
    assert doc.blocks[0].start_char == 0
    assert doc.blocks[0].end_char == 8
    assert doc.blocks[1].start_char == 10
    assert doc.blocks[1].end_char == 15
    assert doc.blocks[1].paragraph == 2
    assert all(b.page is None for b in doc.blocks)
    assert doc.canonical_text_hash == hashlib.sha256(doc.text.encode()).hexdigest()


def test_markdown_sections_and_tables_are_real_and_revision_bound(tmp_path):
    from app.rag.ingestion import load_canonical_document

    doc = load_canonical_document(
        source(
            tmp_path,
            "# 第一章\n定义内容。\n## 条件\n| 条件 | 结果 |\n|---|---|\n| 否 | 不成立 |\n",
            "md",
        )
    )
    assert [s.title for s in doc.sections] == ["第一章", "条件"]
    assert doc.blocks[-1].kind == "table_row"
    assert doc.blocks[-1].heading_path == ["第一章", "条件"]
    assert all(
        s.section_id.startswith(doc.parse_artifact_id + ":") for s in doc.sections
    )
    assert doc.sections[0].end_char == len(doc.text)
    assert doc.blocks[-1].start_char < doc.blocks[-1].end_char


def test_docx_keeps_heading_paragraph_and_table_row_order(tmp_path):
    from docx import Document

    from app.rag.ingestion import load_canonical_document

    request = source(tmp_path, "", "docx")
    raw = Document()
    raw.add_heading("说明", level=1)
    raw.add_paragraph("😀定义。")
    row = raw.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text, row.cells[1].text = "条件", "不得大于 2026"
    raw.add_paragraph("最后一段。")
    raw.save(request.file_path)
    canonical = load_canonical_document(request)
    assert [b.kind for b in canonical.blocks] == [
        "heading",
        "paragraph",
        "table_row",
        "paragraph",
    ]
    assert canonical.text.index("条件") < canonical.text.index("最后一段")
    assert canonical.blocks[2].table_row == 1
    assert canonical.blocks[2].page is None


def test_pdf_uses_one_based_physical_page_and_rejects_blank_scan(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    from app.rag.errors import SourceUnavailable
    from app.rag.ingestion import load_canonical_document

    request = source(tmp_path, "", "pdf")
    writer = PdfWriter()
    for text in ["First page.", "Second page."]:
        page = writer.add_blank_page(300, 300)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): writer._add_object(font)}
                )
            }
        )
        stream = DecodedStreamObject()
        stream.set_data(("BT /F1 12 Tf 20 200 Td (" + text + ") Tj ET").encode())
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(request.file_path)
    canonical = load_canonical_document(request)
    assert [b.page for b in canonical.blocks] == [1, 2]
    assert (
        canonical.text[canonical.blocks[1].start_char : canonical.blocks[1].end_char]
        == "Second page."
    )
    blank = PdfWriter()
    blank.add_blank_page(300, 300)
    blank.write(request.file_path)
    with pytest.raises(SourceUnavailable, match="text|文本"):
        load_canonical_document(request)


def test_parse_id_changes_with_normalizer_and_source_hash(tmp_path):
    from app.rag.ingestion import load_canonical_document

    request = source(tmp_path, "原文")
    original = load_canonical_document(request)
    again = load_canonical_document(request)
    assert original.parse_artifact_id == again.parse_artifact_id
    changed = request.model_copy(update={"source_sha256": "0" * 64})
    from app.rag.errors import SourceUnavailable

    with pytest.raises(SourceUnavailable):
        load_canonical_document(changed)


def test_parser_rejects_zip_bomb_and_oversized_input(tmp_path):
    from app.rag.errors import SourceUnavailable
    from app.rag.ingestion import load_canonical_document

    request = source(tmp_path, "too long")
    with pytest.raises(SourceUnavailable):
        load_canonical_document(request.model_copy(update={"max_bytes": 2}))
    request = source(tmp_path, "", "docx")
    with zipfile.ZipFile(request.file_path, "w", zipfile.ZIP_DEFLATED) as bundle:
        bundle.writestr("word/document.xml", b" " * 1000000)
    with pytest.raises(SourceUnavailable):
        load_canonical_document(request)
